import importlib.util
import os
import shutil
import sys
from pathlib import Path

import pytest
from pptx import Presentation

from scislides import PRESETS, Deck, contrast_ratio, get_theme
from scislides import edit
from scislides.check import check_deck
from scislides.media import montage

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "examples"))


def test_contrast_known_values():
    assert contrast_ratio("000000", "FFFFFF") == pytest.approx(21.0)
    assert contrast_ratio("777777", "777777") == pytest.approx(1.0)


def test_presets_have_enough_contrast():
    assert len(PRESETS) >= 6
    for theme in PRESETS.values():
        assert theme.check() == [], theme.name


def test_unknown_theme():
    with pytest.raises(KeyError):
        get_theme("nope")


@pytest.mark.parametrize("name", list(PRESETS))
def test_demo_builds_and_passes_check(name, tmp_path):
    import demo_deck
    out = demo_deck.build(name, tmp_path / "demo.pptx")
    assert len(Presentation(str(out)).slides) == 8
    errors = [i for i in check_deck(out) if i.severity == "error"]
    assert errors == []


def test_check_finds_problems(tmp_path):
    d = Deck("paper-blue")
    s = d.slide("k", "t")
    sh = d.text(s, "tiny", 1, 1, 2, 1, 20)
    sh.text_frame.paragraphs[0].runs[0].font.size = Pt8()
    from PIL import Image
    Image.new("RGB", (10, 10)).save(tmp_path / "x.png")
    d.pic(s, tmp_path / "x.png", 12, 7, 3, 3)  # off slide, no alt
    path = d.save(tmp_path / "bad.pptx")
    kinds = {i.kind for i in check_deck(path)}
    assert {"font", "offslide", "alt"} <= kinds


def Pt8():
    from pptx.util import Pt
    return Pt(8)


def test_min_size_enforced_and_forbidden():
    d = Deck("paper-blue", forbidden="—")
    s = d.slide()
    sh = d.text(s, "hello", 1, 1, 2, 1, 8)
    assert sh.text_frame.paragraphs[0].runs[0].font.size.pt == d.theme.min_body
    with pytest.raises(ValueError):
        d.text(s, "a — b", 1, 1, 2, 1)


def test_edit_helpers(tmp_path):
    d = Deck("forest-teal")
    for i in range(3):
        s = d.slide("k", f"Title {i}")
        a = d.rect(s, 1, 3, 1, 1, "FF0000", "box")
        d.anim(s, [(a, "fade", "click"), (a, "pulse", "after")])
    path = d.save(tmp_path / "a.pptx")
    prs = edit.open_deck(path)
    edit.set_text(edit.find_shape(prs.slides[0], "title"), "Changed")
    edit.reorder(prs, [2, 0, 1])
    assert edit.renumber(prs) == 3
    first = prs.slides[0]
    assert edit.find_shape(first, "number").text_frame.text == "1"
    edit.remove_shape(edit.find_shape(first, "box"))
    assert edit.drop_orphan_animations(first) == 1
    assert first._element.find("{%s}timing" % edit.NS_P) is None
    out = tmp_path / "b.pptx"
    prs.save(str(out))
    texts = [edit.find_shape(sl, "title").text_frame.text for sl in Presentation(str(out)).slides]
    assert texts == ["Title 2", "Changed", "Title 1"]


def test_replace_blob(tmp_path):
    from PIL import Image
    Image.new("RGB", (20, 20), "red").save(tmp_path / "r.png")
    Image.new("RGB", (20, 20), "blue").save(tmp_path / "b.png")
    d = Deck()
    s = d.slide()
    pic = d.pic(s, tmp_path / "r.png", 1, 1, 2, 2, alt="x")
    edit.replace_blob(pic, tmp_path / "b.png")
    path = d.save(tmp_path / "p.pptx")
    prs = Presentation(str(path))
    blob = [sh for sh in prs.slides[0].shapes if sh.name == "figure"][0].image.blob
    assert blob == (tmp_path / "b.png").read_bytes()


def test_atempo_chain():
    assert montage.atempo_chain(0.33).count("atempo") == 2
    assert montage.atempo_chain(1.0) == "atempo=1.000000"


def test_transition_and_animation_xml_is_wellformed(tmp_path):
    d = Deck()
    s = d.slide("k", "t")
    a = d.rect(s, 1, 3, 1, 1, "FF0000")
    d.anim(s, [(a, "appear", "click"), (a, "fade", "with"), (a, "pulse", "after")])
    d.transition(s)
    d.transition(s, morph=True)
    path = d.save(tmp_path / "m.pptx")
    assert len(Presentation(str(path)).slides) == 1


@pytest.mark.skipif(not (os.environ.get("FFMPEG") or shutil.which("ffmpeg")),
                    reason="ffmpeg not available")
def test_montage_on_synthetic_clip(tmp_path):
    from scislides.media import _ffmpeg
    _ffmpeg.run(["-f", "lavfi", "-i", "testsrc2=s=320x240:r=25:d=3", "-c:v", "libx264",
                 "-pix_fmt", "yuv420p", str(tmp_path / "c.mp4")])
    spec = {"output": "out/m", "size": [320, 240], "fps": 10, "_base": str(tmp_path),
            "segments": [{"source": "c.mp4", "start": 0, "duration": 1, "caption": "hello"},
                         {"source": "c.mp4", "start": 1, "duration": 1, "speed": 0.5,
                          "label": "slow"}],
            "end_card": {"duration": 1, "lines": [{"text": "end"}]},
            "gif": {"width": 160, "fps": 5, "colors": 32}}
    mp4, gif = montage.build_montage(spec)
    assert mp4.stat().st_size > 0 and gif.stat().st_size > 0
    from scislides.media import video
    assert video.contact_sheet(mp4, tmp_path / "sheet.png").exists()


@pytest.mark.skipif(importlib.util.find_spec("rdkit") is None, reason="rdkit not installed")
def test_depict_2d(tmp_path):
    from scislides.media.molecules import depict_2d
    out = depict_2d("CC(=O)Oc1ccccc1C(=O)O", tmp_path / "a.png", background=None)
    from PIL import Image
    assert Image.open(out).mode == "RGBA"
    with pytest.raises(ValueError):
        depict_2d("not a smiles ((", tmp_path / "x.png")


def test_energy_layers(tmp_path):
    from scislides.media.plots import energy_profile_layers
    paths = energy_profile_layers(["A", "B", "C"], [{"label": "x", "values": [0, 1, -1]},
                                                    {"label": "y", "values": [0, 2, -2]}],
                                  tmp_path, "ink-violet", cumulative=True)
    assert len(paths) == 2 and all(p.exists() for p in paths)
