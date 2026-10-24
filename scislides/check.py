"""Deck verifier: catches the problems that are cheap to find and costly on stage."""
from __future__ import annotations

import argparse
import math
import sys
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from .edit import _all_shapes


@dataclass
class Issue:
    slide: int
    shape: str
    kind: str       # font | offslide | overflow | alt | size | char
    message: str
    severity: str = "error"   # error | warning

    def __str__(self) -> str:
        return f"[{self.severity}] slide {self.slide} · {self.shape or '-'} · {self.kind}: {self.message}"


def _estimate_height_in(shape, size_pt: float) -> float:
    """Rough wrapped-text height in inches: average glyph width ~0.5 em, line height 1.2 em."""
    width_in = max(Emu(shape.width).inches, 0.1)
    lines = 0
    for p in shape.text_frame.paragraphs:
        chars = len("".join(r.text for r in p.runs))
        per_line = max(width_in / (size_pt * 0.5 / 72), 1)
        lines += max(1, math.ceil(chars / per_line))
    return lines * size_pt * 1.2 / 72


def check_deck(path, min_font: float = 14, min_number_font: float = 10, max_mb: float = 50,
               forbidden: str = "", overflow_slack: float = 1.15) -> list[Issue]:
    """Return a list of ``Issue``. Page numbers (shape name ``number``) use ``min_number_font``."""
    path = Path(path)
    issues: list[Issue] = []
    mb = path.stat().st_size / 1e6
    if mb > max_mb:
        issues.append(Issue(0, "", "size", f"file is {mb:.1f} MB (budget {max_mb} MB)"))
    prs = Presentation(str(path))
    sw, sh_ = prs.slide_width, prs.slide_height
    for n, slide in enumerate(prs.slides, 1):
        for shp in _all_shapes(slide.shapes):
            name = shp.name
            if shp.shape_type != 6 and None not in (shp.left, shp.top, shp.width, shp.height):
                tol = Emu(9525)  # 1 px
                if shp.left < -tol or shp.top < -tol or shp.left + shp.width > sw + tol or \
                        shp.top + shp.height > sh_ + tol:
                    issues.append(Issue(n, name, "offslide", "shape extends beyond the slide"))
            if shp.shape_type == 13:  # picture
                descr = shp._element.xpath(".//p:cNvPr/@descr")
                if not descr or not descr[0].strip():
                    issues.append(Issue(n, name, "alt", "picture without alt text"))
            if shp.has_text_frame and shp.text_frame.text.strip():
                limit = min_number_font if name == "number" else min_font
                sizes = [r.font.size.pt for p in shp.text_frame.paragraphs for r in p.runs
                         if r.font.size is not None]
                if sizes and min(sizes) < limit:
                    issues.append(Issue(n, name, "font", f"{min(sizes):g} pt < {limit:g} pt"))
                if sizes and shp.shape_type != 6:
                    need = _estimate_height_in(shp, max(sizes))
                    have = Emu(shp.height).inches
                    if need > have * overflow_slack and need > 0.35:
                        issues.append(Issue(n, name, "overflow",
                                            f"text needs ~{need:.2f} in, box is {have:.2f} in",
                                            "warning"))
                bad = [c for c in forbidden if c in shp.text_frame.text]
                if bad:
                    issues.append(Issue(n, name, "char", f"forbidden character(s) {bad!r}"))
    return issues


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="scislides-check", description=__doc__)
    ap.add_argument("deck", nargs="+")
    ap.add_argument("--min-font", type=float, default=14)
    ap.add_argument("--max-mb", type=float, default=50)
    ap.add_argument("--forbid", default="", help="characters that must not appear, e.g. '—'")
    ap.add_argument("--strict", action="store_true", help="treat warnings as errors")
    a = ap.parse_args(argv)
    bad = 0
    for deck in a.deck:
        issues = check_deck(deck, a.min_font, max_mb=a.max_mb, forbidden=a.forbid)
        print(f"{deck}: {len(issues)} issue(s)")
        for i in issues:
            print("  ", i)
            bad += i.severity == "error" or a.strict
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
