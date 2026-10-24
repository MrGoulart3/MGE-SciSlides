"""``Deck``: a thin, opinionated layer over python-pptx for 16:9 scientific slides.

All positions are in inches on a 13.333 x 7.5 grid with a 0.75 in side margin.
Shape names are meaningful (``kicker``, ``title``, ``number``, ``caption`` ...):
the checker and the edit helpers rely on them.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from . import motion
from .theme import Theme, get_theme

W, H = 13.333333, 7.5
MARGIN = 0.75
CONTENT_W = W - 2 * MARGIN
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
# Shape-name prefixes that count as small text (captions, credits, sources, kickers).
CAPTION_PREFIXES = ("caption", "credit", "source", "kicker", "label")


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#"))


class Deck:
    """A presentation under construction.

    ``forbidden`` is an optional string of characters that must not appear in text
    (e.g. ``"—"`` if your style guide bans em dashes); ``text`` raises on them.
    """

    def __init__(self, theme: "str | Theme | None" = None, title: str = "",
                 author: str = "", subject: str = "", lang: str = "en-US",
                 forbidden: str = "") -> None:
        self.theme = get_theme(theme)
        self.lang = lang
        self.forbidden = forbidden
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        cp = self.prs.core_properties
        cp.title, cp.author, cp.subject, cp.language = title, author, subject, lang
        cp.last_modified_by = author
        cp.created = cp.modified = dt.datetime.now().replace(microsecond=0)

    # ----------------------------------------------------------------- slides
    def slide(self, kicker: str | None = None, title: str | None = None,
              dark: bool = False, size: int = 32):
        """Add a blank slide with optional kicker, assertion title and page number."""
        t = self.theme
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = t.dark_bg if dark else t.background
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg)
        ink = t.dark_ink if dark else t.ink
        if kicker:
            self.text(s, kicker.upper(), MARGIN, 0.6, CONTENT_W, 0.30, t.min_caption,
                      t.on_dark_accent if dark else t.accent, t.font_title, "kicker",
                      tracking=100)
        if title:
            self.text(s, title, MARGIN, 1.11, CONTENT_W, 1.14, size, ink, t.font_title,
                      "title", spacing=0.95)
        self.text(s, str(len(self.prs.slides)), 12.20, 6.74, 0.38, 0.20, t.min_number,
                  t.dark_ink if dark else t.muted, t.font_body, "number", align="right")
        return s

    # ------------------------------------------------------------------ shapes
    def text(self, s, text: str, x: float, y: float, w: float, h: float, size: int = 20,
             color: str | None = None, font: str | None = None, name: str = "text",
             align: str | None = None, spacing: float = 1.05, tracking: int | None = None,
             bold: bool = False):
        """Add a text box. Sizes below the theme minimum are raised, never lowered."""
        t = self.theme
        bad = [c for c in self.forbidden if c in text]
        if bad:
            raise ValueError(f"forbidden character(s) {bad!r} in {text!r}")
        if name == "number":
            floor = t.min_number
        elif name.startswith(CAPTION_PREFIXES):
            floor = t.min_caption
        else:
            floor = t.min_body
        size = max(size, floor)
        sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        sh.name = name
        tf = sh.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.line_spacing, p.space_before, p.space_after = spacing, Pt(0), Pt(0)
            if align:
                p.alignment = ALIGN[align]
            for r in p.runs:
                r.font.name = font or t.font_body
                r.font.size = Pt(size)
                r.font.bold = bold or None
                r.font.color.rgb = rgb(color or t.ink)
                rpr = r._r.get_or_add_rPr()
                rpr.set("lang", self.lang)
                if tracking:
                    rpr.set("spc", str(tracking))
        return sh

    def rect(self, s, x, y, w, h, color, name="shape", alpha: float | None = None):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.name = name
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(color)
        sh.line.fill.background()
        if alpha is not None:
            node = sh.fill._xPr.solidFill[0]
            node.append(OxmlElement("a:alpha"))
            node[-1].set("val", str(round(alpha * 100000)))
        return sh

    def line(self, s, x1, y1, x2, y2, color=None, width=0.75, name="rule", dashed=False):
        sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
        sh.name = name
        sh.line.color.rgb = rgb(color or self.theme.rule)
        sh.line.width = Pt(width)
        if dashed:
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return sh

    def circle(self, s, x, y, d, color=None, name="dot", outline=None):
        sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        sh.name = name
        if color:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(color)
        else:
            sh.fill.background()
        if outline:
            sh.line.color.rgb = rgb(outline)
            sh.line.width = Pt(1.5)
        else:
            sh.line.fill.background()
        return sh

    def pic(self, s, path, x, y, w, h, alt: str = "", name: str = "figure", fit: bool = False):
        """Place a picture in the box. ``fit=True`` letterboxes; otherwise center-crop to fill."""
        path = Path(path)
        with Image.open(path) as im:
            iw, ih = im.size
        if fit:
            k = min(w / iw, h / ih)
            rw, rh = iw * k, ih * k
            sh = s.shapes.add_picture(str(path), Inches(x + (w - rw) / 2),
                                      Inches(y + (h - rh) / 2), Inches(rw), Inches(rh))
        else:
            sh = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
            ratio, target = iw / ih, w / h
            if ratio > target:
                sh.crop_left = sh.crop_right = (1 - target / ratio) / 2
            else:
                sh.crop_top = sh.crop_bottom = (1 - ratio / target) / 2
        sh.name = name
        sh._element.xpath(".//p:cNvPr")[0].set("descr", alt)
        return sh

    def full_bleed(self, s, path, alt: str = "", name: str = "background_image"):
        """Cover the whole slide with a picture and send it behind every other shape."""
        sh = self.pic(s, path, 0, 0, W, H, alt=alt, name=name)
        tree = s.shapes._spTree
        tree.remove(sh._element)
        tree.insert(2, sh._element)  # right after nvGrpSpPr and grpSpPr
        return sh

    def notes(self, s, text: str) -> None:
        """Speaker notes."""
        s.notes_slide.notes_text_frame.text = text

    def group(self, s, shapes, name: str = "group"):
        g = s.shapes.add_group_shape(shapes)
        g.name = name
        return g

    # ------------------------------------------------------------- motion/flow
    def anim(self, s, rows, duration: float = 0.4) -> None:
        """Click animations. ``rows`` = ``[(shape, effect, trigger[, options])]``.

        effect: ``appear`` | ``fade`` | ``pulse``; trigger: ``click`` | ``with`` | ``after``.
        """
        spec = []
        for sh, effect, trigger, *opt in rows:
            row = {"shape_id": sh.shape_id, "effect": effect, "trigger": trigger,
                   "is_sp": sh._element.tag.endswith("}sp")}
            if opt:
                row.update(opt[0])
            spec.append(row)
        motion.set_timing(s._element, spec, duration)

    def transition(self, s, morph: bool = False) -> None:
        """Fade (default) or morph transition into this slide."""
        motion.set_transition(s._element, "morph" if morph else "fade")

    def carrier(self, s, index: int, total: int, labels: list[str] | None = None,
                y: float = 0.4, dark: bool = False, end: float | None = None):
        """Thin progress line with ``total`` stations; the marker sits at station ``index`` (1-based)."""
        t = self.theme
        accent = t.on_dark_accent if dark else t.accent
        x2 = end if end is not None else W - MARGIN
        pos = MARGIN + (x2 - MARGIN) * index / total
        pos = min(pos, x2 - 0.055)
        self.line(s, MARGIN, y, x2, y, t.rule, 0.75, "progress_track")
        self.line(s, MARGIN, y, pos, y, accent, 1.25, "progress_done")
        dot = self.circle(s, pos - 0.055, y - 0.055, 0.11, accent, "progress_marker")
        label = labels[index - 1] if labels and 0 < index <= len(labels) else f"{index}/{total}"
        dot._element.xpath(".//p:cNvPr")[0].set("descr", f"Progress: {label}")
        return dot

    def save(self, path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path
